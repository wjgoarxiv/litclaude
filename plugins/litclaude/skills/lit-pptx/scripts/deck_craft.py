#!/usr/bin/env python3
"""deck_craft.py — measurable deck-craft checks for the lit-pptx QA gate.

The layout inventory answers "does anything spill or collide". These checks answer
"does the deck look finished", with numbers a gate can hold:

  sparse_slide        a content slide whose content fills less than FILL_MIN of the
                      area under its title (a short table at the top of an empty slide)
  table_only_deck     three or more content slides, tables on at least 60% of them,
                      and not one chart in the deck
  off_slide_shape     a picture or shape that runs past the canvas edge and is not a
                      full-canvas background
  stray_box           a large outlined rectangle with no fill and no text in or on it
  placeholder_text    unfilled blanks: [제품명], [금액], XXX, TBD, ○○, {{…}}
  long_foreign_note   on a Korean deck, a source line (or any small caption) that is
                      mostly a long run of Latin text

Each finding carries the slide, the shape and a fix hint. A slide counts as a
display slide (cover, section, closing) when it holds text at 30 pt or more; display
slides are exempt from the fill and table checks.

Usage:
    python3 deck_craft.py deck.pptx          # JSON report; exit 0 always
"""

from __future__ import annotations

import os as _os
import sys as _sys

_sys.dont_write_bytecode = True  # the installed skill directory stays read-only
_sys.path.insert(0, _os.path.join(_os.path.dirname(_os.path.abspath(__file__)), "..", "..", "..", "lib"))
from office_runtime_bootstrap import ensure_runtime  # noqa: E402

ensure_runtime(["pptx"])

import json  # noqa: E402
import re  # noqa: E402
from pathlib import Path  # noqa: E402

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402
from pptx.util import Emu  # noqa: E402

EMU = 914400.0
FILL_MIN = 0.45
DISPLAY_PT = 30.0
CAPTION_PT = 10.5
EDGE_TOL = 0.03
FULL_BLEED = 0.92
# A bracket with a digit in it is a citation marker ([1], [S3]); a blank has none.
PLACEHOLDER = re.compile(
    r"\[(?![^\]\n]*\d)[^\]\n]{1,24}\]|\bX{3,}\b|\bTBD\b|\bTODO\b|○○|OOO|\{\{[^}]*\}\}|_{4,}|＿{2,}|<[^<>\n]{1,20}(?:명|값|금액|날짜|이름)>"
)
MARKDOWN_LINK = re.compile(r"\[[^\]]+\]\([^)]+\)")
HANGUL = re.compile(r"[가-힣]")
LATIN = re.compile(r"[A-Za-z]")


def _box(shape):
    """(left, top, width, height) in inches; a table's height is the sum of its rows,
    which renderers grow to, not the frame height a writer declared."""
    try:
        if shape.left is None or shape.width is None:
            return None
        height = shape.height / EMU
        if getattr(shape, "has_table", False) and shape.has_table:
            height = max(height, sum(row.height for row in shape.table.rows) / EMU)
        return (shape.left / EMU, shape.top / EMU, shape.width / EMU, height)
    except (ValueError, TypeError):
        # Writers sometimes emit fractional EMU values python-pptx refuses to parse.
        return None


def _flatten(shapes):
    """Top-level shapes only; a group counts as one box. Group children sit in the
    group's own coordinate space, which is not the slide's."""
    for shape in shapes:
        yield shape


def _max_pt(shape):
    """Largest font size in the shape: run sizes, then paragraph defaults and end-of-paragraph
    sizes, then PowerPoint's 18 pt default when a writer set none at all."""
    if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
        return 0.0
    sizes = []
    for element in shape.text_frame._txBody.iter():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag in ("rPr", "defRPr", "endParaRPr") and element.get("sz"):
            try:
                sizes.append(int(element.get("sz")) / 100.0)
            except ValueError:
                pass
    return max(sizes) if sizes else 18.0


def _text(shape):
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        return shape.text_frame.text.strip()
    if getattr(shape, "has_table", False) and shape.has_table:
        return "\n".join(cell.text for row in shape.table.rows for cell in row.cells)
    return ""


def _is_chart(shape):
    return getattr(shape, "has_chart", False) and shape.has_chart


def _is_table(shape):
    return getattr(shape, "has_table", False) and shape.has_table


def _descr(shape):
    element = shape._element
    for tag in element.iter():
        if tag.tag.endswith("}cNvPr"):
            return (tag.get("descr") or "") + " " + (tag.get("name") or "")
    return ""


def _decorative_picture(shape, page_area):
    box = _box(shape)
    if box is None:
        return True
    if box[2] * box[3] >= FULL_BLEED * page_area:
        return True
    return bool(re.search(r"circle-|template-image|deco", _descr(shape), re.I))


def _filled(shape):
    try:
        return shape.fill.type is not None and shape.fill.type != 5  # 5 = background
    except (AttributeError, TypeError):
        return False


def _outlined(shape):
    line = shape._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}ln")
    if line is None:
        return False
    if line.find("{http://schemas.openxmlformats.org/drawingml/2006/main}noFill") is not None:
        return False
    return len(line) > 0 or line.get("w") is not None


def _contains(outer, inner, tol=0.05):
    return (inner[0] >= outer[0] - tol and inner[1] >= outer[1] - tol
            and inner[0] + inner[2] <= outer[0] + outer[2] + tol and inner[1] + inner[3] <= outer[1] + outer[3] + tol)


def slide_layout(shapes, W, H, page_area):
    """The slide's display flag, title, content area and content boxes, shared by the craft
    checks here and in craft_extras.py so both judge the same region."""
    texts = [(s, _text(s)) for s in shapes]
    # Cover, section and closing slides: display-size type, no data visual, and at
    # most two headline-sized lines. A KPI slide has several large numbers and is content.
    has_visual = any(_is_table(s) or _is_chart(s) or (s.shape_type == MSO_SHAPE_TYPE.PICTURE and not _decorative_picture(s, page_area))
                     for s in shapes)
    large = [s for s, t in texts if t and _max_pt(s) >= 18]
    display = (any(_max_pt(s) >= DISPLAY_PT for s, t in texts if t) and not has_visual and len(large) <= 2)

    # Title: the largest text near the top.
    tops = [(s, _max_pt(s)) for s, t in texts if t and _box(s) and _box(s)[1] < H * 0.25]
    title = max(tops, key=lambda item: item[1])[0] if tops else None
    title_bottom = (_box(title)[1] + _box(title)[3]) if title is not None else H * 0.2
    area = (0.5, title_bottom + 0.1, W - 1.0, max(0.1, H - 0.55 - title_bottom - 0.1))

    visuals = []
    boxes = []
    for shape, text in texts:
        if shape is title:
            continue
        box = _box(shape)
        if box is None:
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            if _decorative_picture(shape, page_area):
                continue
            visuals.append("figure")
        elif _is_chart(shape):
            visuals.append("chart")
        elif _is_table(shape):
            visuals.append("table")
        elif text:
            if _max_pt(shape) <= CAPTION_PT:
                continue
        elif shape.shape_type in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.GROUP) and box[2] * box[3] >= 0.5 and (
                shape.shape_type == MSO_SHAPE_TYPE.GROUP or _filled(shape)):
            if box[2] * box[3] < FULL_BLEED * page_area:
                visuals.append("card")
        else:
            continue
        boxes.append(box)
    return {"display": display, "title": title, "area": area, "boxes": boxes, "visuals": visuals}


def analyse(path: Path) -> dict:
    prs = Presentation(str(path))
    W, H = prs.slide_width / EMU, prs.slide_height / EMU
    page_area = W * H
    findings: list[dict] = []
    all_text = []
    per_slide = []

    for number, slide in enumerate(prs.slides, start=1):
        # The engine's sample-data tag is a label on every slide, not slide content.
        shapes = [s for s in _flatten(slide.shapes) if not s.name.startswith("lit-notice")]
        texts = [(s, _text(s)) for s in shapes]
        all_text.extend(t for _, t in texts if t)
        layout = slide_layout(shapes, W, H, page_area)
        display = layout["display"]

        # Placeholder blanks anywhere, display slides included.
        for shape, text in texts:
            cleaned = MARKDOWN_LINK.sub(" ", text)
            hit = PLACEHOLDER.search(cleaned)
            if hit:
                findings.append({"check": "placeholder_text", "slide": number, "shape": shape.name,
                                 "detail": hit.group(0),
                                 "hint": "Fill the blank with a realistic value and label it as an example or assumption."})

        # Anything past the canvas edge that is not a full-canvas background.
        for shape in shapes:
            box = _box(shape)
            if box is None or box[2] * box[3] >= FULL_BLEED * page_area:
                continue
            if box[0] < -EDGE_TOL or box[1] < -EDGE_TOL or box[0] + box[2] > W + EDGE_TOL or box[1] + box[3] > H + EDGE_TOL:
                findings.append({"check": "off_slide_shape", "slide": number, "shape": shape.name,
                                 "detail": [round(v, 2) for v in box],
                                 "hint": "Keep decorations inside the canvas; a cropped shape reads as a layout error."})

        # Large outlined empty rectangles.
        content_boxes = [_box(s) for s, t in texts if (t or _is_chart(s) or s.shape_type == MSO_SHAPE_TYPE.PICTURE) and _box(s)]
        for shape in shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE or _text(shape):
                continue
            box = _box(shape)
            if box is None or box[2] * box[3] < 1.0:
                continue
            try:
                geometry = shape.auto_shape_type
                rectangular = "RECTANGLE" in str(geometry)
            except (NotImplementedError, ValueError):
                rectangular = False
            if not rectangular or _filled(shape) or not _outlined(shape):
                continue
            if any(_contains(box, inner) for inner in content_boxes if inner != box):
                continue
            findings.append({"check": "stray_box", "slide": number, "shape": shape.name,
                             "detail": [round(v, 2) for v in box],
                             "hint": "Remove the empty frame or put the message inside a tinted callout."})

        if display:
            per_slide.append({"slide": number, "display": True})
            continue

        title, area, boxes, visuals = layout["title"], layout["area"], layout["boxes"], layout["visuals"]

        if not boxes and title is not None and len(_text(title)) <= 30:
            # Only a short line on the slide: a closing or "Q&A" slide, not a sparse one.
            per_slide.append({"slide": number, "display": True})
            continue
        if boxes:
            left = max(area[0], min(b[0] for b in boxes))
            top = max(area[1], min(b[1] for b in boxes))
            right = min(area[0] + area[2], max(b[0] + b[2] for b in boxes))
            bottom = min(area[1] + area[3], max(b[1] + b[3] for b in boxes))
            fill = max(0.0, right - left) * max(0.0, bottom - top) / (area[2] * area[3])
        else:
            fill = 0.0
        cards = visuals.count("card")
        primary = ("chart" if "chart" in visuals else "figure" if "figure" in visuals else "table" if "table" in visuals
                   else "cards" if cards >= 2 else "text")
        per_slide.append({"slide": number, "display": False, "fill": round(fill, 2), "primary": primary})
        if fill < FILL_MIN:
            findings.append({"check": "sparse_slide", "slide": number, "shape": None, "detail": round(fill, 2),
                             "hint": f"Content fills {fill:.0%} of the area under the title (floor {FILL_MIN:.0%}): "
                                     "give the message a visual that uses the space (chart, KPI cards, a fuller table "
                                     "with its takeaway), or merge the slide with its neighbour."})

    content = [s for s in per_slide if not s["display"]]
    tables = sum(1 for s in content if s["primary"] == "table")
    charts = sum(1 for s in content if s["primary"] == "chart")
    if len(content) >= 3 and charts == 0 and tables >= 2 and tables >= 0.6 * len(content):
        findings.append({"check": "table_only_deck", "slide": None, "shape": None,
                         "detail": f"{tables} of {len(content)} content slides are tables and none is a chart",
                         "hint": "Numeric series belong in native charts (::: chart) and headline numbers in KPI "
                                 "cards; keep tables for lookups a reader scans."})

    corpus = "\n".join(all_text)
    # One Hangul syllable carries about as much as two or three Latin letters.
    korean = 2.5 * len(HANGUL.findall(corpus)) > len(LATIN.findall(corpus))
    if korean:
        for number, slide in enumerate(prs.slides, start=1):
            for shape in _flatten(slide.shapes):
                text = _text(shape)
                if not text:
                    continue
                source_line = bool(re.match(r"\s*(?:sources?|출처|자료)\b", text, re.I))
                if _max_pt(shape) > CAPTION_PT + 0.5 and not source_line:
                    continue
                latin, hangul = len(LATIN.findall(text)), len(HANGUL.findall(text))
                if latin >= 60 and latin / max(1, latin + hangul) > 0.7:
                    findings.append({"check": "long_foreign_note", "slide": number, "shape": shape.name,
                                     "detail": text[:80],
                                     "hint": "On a Korean deck keep sources short and in Korean (기관, 문서명, 연도)."})

    return {"file": str(path), "pass": not findings, "slides": per_slide, "findings": findings}


def main(argv: list[str]) -> int:
    if not argv:
        print(__doc__)
        return 2
    print(json.dumps(analyse(Path(argv[0])), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(_sys.argv[1:]))
