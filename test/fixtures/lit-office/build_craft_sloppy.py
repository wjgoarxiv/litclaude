"""Build a deck that plants each office craft defect (OF-101..OF-109) for the lit-pptx gate test.

Usage: python build_craft_sloppy.py out.pptx
"""
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def box(slide, x, y, w, h, fill=None, shape=MSO_SHAPE.RECTANGLE, name=None):
    item = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        item.fill.solid()
        item.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        item.fill.background()
    item.line.fill.background()
    if name:
        item.name = name
    return item


def text(slide, x, y, w, h, value, size=16, name=None):
    frame = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame.text_frame.word_wrap = True
    frame.text_frame.text = value
    for paragraph in frame.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size)
    if name:
        frame.name = name
    return frame


def title(slide, value):
    return text(slide, 0.5, 0.3, 9, 0.8, value, size=28, name="Title")


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]

# 1. Three accent families (OF-101), gradient text (OF-106), glow (OF-107), emoji bullet (OF-108).
slide = prs.slides.add_slide(blank)
title(slide, "Quarter overview")
for x, fill in ((0.5, "D62828"), (4.6, "2A9D8F"), (8.7, "7B2CBF")):
    box(slide, x, 1.6, 3.8, 1.2, fill=fill)
gradient = text(slide, 0.5, 3.2, 12, 0.8, "Shiny headline", size=24, name="Gradient")
rpr = gradient.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr()
grad = etree.SubElement(rpr, qn("a:gradFill"))
stops = etree.SubElement(grad, qn("a:gsLst"))
for pos, color in (("0", "FF0080"), ("100000", "7928CA")):
    stop = etree.SubElement(stops, qn("a:gs"), pos=pos)
    etree.SubElement(stop, qn("a:srgbClr"), val=color)
glowing = box(slide, 0.5, 4.3, 3, 1, fill="EEEEEE", name="Glowing")
effects = etree.SubElement(glowing._element.spPr, qn("a:effectLst"))
glow = etree.SubElement(effects, qn("a:glow"), rad="228600")
etree.SubElement(glow, qn("a:srgbClr"), val="FFD700")
bullets = text(slide, 4.6, 4.3, 8, 1.5, "Point one\nPoint two", size=16, name="Bullets")
for paragraph in bullets.text_frame.paragraphs:
    ppr = paragraph._p.get_or_add_pPr()
    etree.SubElement(ppr, qn("a:buChar"), char="\U0001F680")

# 2. A long wrapped body line (OF-102), a left-aligned numeric column (OF-103), an empty band (OF-109).
slide = prs.slides.add_slide(blank)
title(slide, "Details")
text(slide, 0.5, 1.4, 12.3, 1.2, "This body paragraph runs across the whole slide width so each wrapped line carries far more characters than anyone can comfortably read in one pass at presentation distance, which is what the measure check exists to catch.", size=14, name="Wide body")
rows, cols = 4, 3
table = slide.shapes.add_table(rows, cols, Inches(0.5), Inches(2.9), Inches(8), Inches(1.6)).table
for c, head in enumerate(["Item", "Q1", "Q2"]):
    table.cell(0, c).text = head
for r, (item, q1, q2) in enumerate([("North", "1,184", "1,220"), ("South", "101", "96"), ("East", "78", "81")], start=1):
    for c, value in enumerate((item, q1, q2)):
        cell = table.cell(r, c)
        cell.text = value
        cell.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT

# 3. A symmetric frame whose inner radius ignores the padding (OF-104), cards closer than twice the gap between their contents (OF-105).
slide = prs.slides.add_slide(blank)
title(slide, "Frames and cards")
outer = box(slide, 0.5, 1.5, 4, 3, fill="E8EEF7", shape=MSO_SHAPE.ROUNDED_RECTANGLE, name="Outer frame")
outer.adjustments[0] = 0.2
inner = box(slide, 0.9, 1.9, 3.2, 2.2, fill="FFFFFF", shape=MSO_SHAPE.ROUNDED_RECTANGLE, name="Inner frame")
inner.adjustments[0] = 0.3
for x in (5.0, 7.7, 10.4):
    box(slide, x, 1.5, 2.6, 2.6, fill="F2F2F2", name="Card")
    text(slide, x + 0.4, 1.9, 1.8, 0.6, "Card", size=16)
    text(slide, x + 0.4, 2.8, 1.8, 0.6, "Detail", size=14)

prs.save(sys.argv[1])
