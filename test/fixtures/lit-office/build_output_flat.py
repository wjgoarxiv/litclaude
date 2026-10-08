"""Build a deck that fails the deck output checks (OF-110..OF-113) for the lit-pptx gate test.

It is shaped like a deck the engine built (the core-properties stamp, `family@` and `title@`
shape names), but every content slide uses one title treatment over one short column, which is
what the engine produced before tonality packs: one look repeated, the lower half of each page
empty. Slide 5 nudges its title 12 pt off the frame and slide 7 names a treatment it does not draw.

Usage: python build_output_flat.py out.pptx
"""
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt


def rect(slide, x, y, w, h, fill, name=None):
    item = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    item.fill.solid()
    item.fill.fore_color.rgb = RGBColor.from_string(fill)
    item.line.fill.background()
    if name:
        item.name = name
    return item


def text(slide, x, y, w, h, value, size, name=None, bold=False):
    frame = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    frame.text_frame.word_wrap = True
    frame.text_frame.text = value
    for paragraph in frame.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string("16212C")
    if name:
        frame.name = name
    return frame


prs = Presentation()
prs.slide_width, prs.slide_height = Pt(960), Pt(540)
prs.core_properties.subject = "lit-pptx tonality=ledger density=6 grid=standard variance=5"
blank = prs.slide_layouts[6]

cover = prs.slides.add_slide(blank)
rect(cover, 48, 132, 50, 4, "0E6B5A", name="family@cover-typographic")
text(cover, 48, 144, 716, 120, "3분기 운영 점검", 44, name="title@cover", bold=True)

titles = ["처리량은 목표를 넘었다", "지연은 판독 구간에서 생긴다", "반품 혼입이 재분류를 만든다", "판독기 교체가 가장 효과가 크다",
          "라벨 재인쇄는 비용이 작다", "야간 조는 처리량이 낮다", "교체 일정은 11월 말까지다", "승인을 요청드립니다"]
for number, title in enumerate(titles, start=2):
    slide = prs.slides.add_slide(blank)
    rect(slide, 48, 102, 50, 4, "0E6B5A", name="family@text-column")
    x = 60 if number == 5 else 48
    name = "title@side-rail" if number == 7 else "title@top-rule"
    text(slide, x, 36, 864, 60, title, 28, name=name, bold=True)
    text(slide, 48, 120, 642, 90, "- 근거 한 줄\n- 근거 두 줄", 18)

prs.save(sys.argv[1])
